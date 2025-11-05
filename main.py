import sys
from pathlib import Path

from interface import Ui_MainWindow
from PyQt5 import QtCore, QtWidgets
from PyQt5.QtCore import QUrl
from PyQt5.QtWidgets import QFileDialog
from PyQt5.QtMultimedia import QMediaContent, QMediaPlayer
from docx import Document
from pydub import AudioSegment

import whisper


class ProgressHandler(QtCore.QThread):
    mySignal = QtCore.pyqtSignal(list)

    def __init__(self, input_file: Path, output_folder: Path, output_text_file: Path, output_docx_file: Path, parent=None):
        super().__init__(parent)
        self.input_file = Path(input_file)
        self.output_folder = Path(output_folder)
        self.output_text_file = Path(output_text_file)
        self.output_docx_file = Path(output_docx_file)

    def run(self):
        try:
            transcripts = []

            def split_audio(input_file: Path, output_folder: Path, segment_length: int = 30) -> int:
                output_folder.mkdir(parents=True, exist_ok=True)
                audio_file = AudioSegment.from_file(str(input_file))
                segment_length_ms = segment_length * 1000
                count = 0

                for count, start_time in enumerate(range(0, len(audio_file), segment_length_ms), start=1):
                    segment = audio_file[start_time:start_time + segment_length_ms]
                    output_file = output_folder / f"segment_{count}.wav"
                    segment.export(str(output_file), format="wav")

                return count if count else 0

            segments_count = split_audio(self.input_file, self.output_folder)

            if segments_count == 0:
                raise ValueError("Не удалось разбить аудио на сегменты")

            model = whisper.load_model("large-v3")

            for step in range(segments_count):
                audio_path = self.output_folder / f"segment_{step + 1}.wav"
                audio = whisper.load_audio(str(audio_path))
                audio = whisper.pad_or_trim(audio)

                mel = whisper.log_mel_spectrogram(audio, n_mels=128).to(model.device)
                options = whisper.DecodingOptions()

                result = whisper.decode(model, mel, options)
                transcripts.append(result.text.strip())
                progress_value = ((step + 1) * 100) // segments_count
                self.mySignal.emit(["progress_increment", progress_value])

            final_text = "\n".join(filter(None, transcripts)).strip()

            with self.output_text_file.open("w", encoding="utf-8") as result_file:
                result_file.write(final_text)

            document = Document()
            document.add_heading("Отчёт по транскрибации", level=1)
            info_paragraph = document.add_paragraph()
            info_run = info_paragraph.add_run(f"Исходный файл: {self.input_file.name}\n")
            info_run.bold = True
            document.add_paragraph(f"Количество сегментов: {segments_count}")
            document.add_paragraph("")
            document.add_heading("Текст транскрибации", level=2)
            if final_text:
                for block in final_text.splitlines():
                    if block.strip():
                        document.add_paragraph(block)
            else:
                document.add_paragraph("Речь не была распознана либо файл не содержит звуковых данных.")
            document.save(str(self.output_docx_file))

            self.mySignal.emit([
                "transcription_finished",
                final_text,
                str(self.output_text_file),
                str(self.output_docx_file),
            ])
        except Exception as exc:  # pragma: no cover - safety net for UI thread
            self.mySignal.emit(["transcription_failed", str(exc)])


class MainWindow(QtWidgets.QMainWindow):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.ui = Ui_MainWindow()
        self.ui.setupUi(self)

        self.media_player = QMediaPlayer()
        self.ui.progressBar.setValue(0)

        self.ui.choose_file_bt.clicked.connect(self.transcribe_audio)

    def transcribe_audio(self):
        file_path, _ = QFileDialog.getOpenFileName(
            self,
            "Выберите аудиофайл",
            "",
            "Audio Files (*.wav *.mp3 *.flac *.m4a *.ogg)"
        )

        if not file_path:
            return

        input_path = Path(file_path)
        self.ui.file_name_value.setText(input_path.name)
        self.ui.status_label.setText("Файл загружен, начинаем обработку…")
        self.ui.transcription_preview.clear()
        self.ui.progressBar.setValue(0)
        self.ui.choose_file_bt.setEnabled(False)

        output_folder_path = input_path.parent / "output_segments"
        output_text_file = input_path.with_name(f"{input_path.stem}_transcript.txt")
        output_docx_file = input_path.with_name(f"{input_path.stem}_report.docx")

        self.transcriber_thread = ProgressHandler(
            input_path,
            output_folder_path,
            output_text_file,
            output_docx_file,
        )
        self.transcriber_thread.mySignal.connect(self.signal_handler)
        self.transcriber_thread.start()

    def signal_handler(self, value):
        event_type = value[0]

        if event_type == "progress_increment":
            progress_value = value[1]
            self.ui.progressBar.setValue(progress_value)
            self.ui.status_label.setText(f"Обработка сегментов… {progress_value}%")

        elif event_type == "transcription_finished":
            _, final_text, txt_path, docx_path = value
            self.ui.progressBar.setValue(100)
            self.ui.status_label.setText("Готово! Отчёт сохранён.")
            self.ui.transcription_preview.setPlainText(final_text if final_text else "Речь не была распознана.")
            self.ui.choose_file_bt.setEnabled(True)
            self.play_completion_sound()
            self.show_completion_toast(txt_path, docx_path)

        elif event_type == "transcription_failed":
            _, error_message = value
            self.ui.status_label.setText(f"Ошибка: {error_message}")
            self.ui.choose_file_bt.setEnabled(True)
            QtWidgets.QMessageBox.critical(self, "Ошибка", error_message)

    def play_completion_sound(self):
        sound_path = Path(__file__).resolve().parent / "sound" / "1.mp3"
        if sound_path.exists():
            media_content = QMediaContent(QUrl.fromLocalFile(str(sound_path)))
            self.media_player.setMedia(media_content)
            self.media_player.play()

    def show_completion_toast(self, txt_path: str, docx_path: str):
        message = (
            f"Готово!\nTXT: {txt_path}\nDOCX: {docx_path}"
        )
        QtWidgets.QMessageBox.information(self, "Транскрибация завершена", message)


if __name__ == "__main__":
    app = QtWidgets.QApplication(sys.argv)
    window = MainWindow()
    window.show()
    sys.exit(app.exec_())
